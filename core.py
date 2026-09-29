"""
Lógica de negocio del generador de reportes de ventas.
Compartida entre la versión de escritorio (Tkinter) y, si hiciera falta, la web (Streamlit).
"""
import os
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.chart import PieChart, BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.trendline import Trendline

COL_TOTAL = 'Total Total'
COL_SUBTOTAL = 'Neto Gravado Neto Gravado'
COL_IVA = 'IVA IVA'

COLUMNAS_REGLA = ['Vendedor', 'Usuario Creador', 'Localidad Documento', 'Sucursal']
COLUMNAS_CSV_REGLAS = [
    'Sucursal Origen (Opcional)', 'Columna a Buscar', 'Contiene el Texto',
    'Asignar a Sucursal', 'Asignar a Razón Social'
]


def formatear_mes_es(fecha_str):
    if pd.isna(fecha_str) or str(fecha_str).lower() == 'nat':
        return "Mes_Desconocido"
    meses = {'01': 'Enero', '02': 'Febrero', '03': 'Marzo', '04': 'Abril', '05': 'Mayo', '06': 'Junio',
             '07': 'Julio', '08': 'Agosto', '09': 'Septiembre', '10': 'Octubre', '11': 'Noviembre', '12': 'Diciembre'}
    try:
        y, m = str(fecha_str).split('-')
        return f"{meses[m]} {y}"
    except Exception:
        return str(fecha_str)


def reglas_por_defecto():
    return pd.DataFrame({
        'Sucursal Origen (Opcional)': ['PAVON', 'PAVON', 'PAVON', 'PAVON', 'PAVON', ''],
        'Columna a Buscar': ['Vendedor', 'Usuario Creador', 'Localidad Documento', 'Usuario Creador', 'Localidad Documento', 'Usuario Creador'],
        'Contiene el Texto': ['Victor Bascuñan', 'Caniuqueo Veronica', '00014-', 'Roman Ayrton', '00020-', 'Aguilar Mónica'],
        'Asignar a Sucursal': ['Geo Mayorista', 'Dunia Mayorista', 'Mercado Libre', 'Mercado Libre', 'Mercado Libre', 'EXCLUIR'],
        'Asignar a Razón Social': ['Mnos Import SRL', 'Dunia Import SRL', 'Bekir 1 SRL', 'Bekir 1 SRL', 'Bekir 1 SRL', '']
    })


def cargar_reglas(archivo_reglas):
    if os.path.exists(archivo_reglas):
        df = pd.read_csv(archivo_reglas).fillna('')
        for col in COLUMNAS_CSV_REGLAS:
            if col not in df.columns:
                df[col] = ''
        return df[COLUMNAS_CSV_REGLAS]
    return reglas_por_defecto()


def guardar_reglas(df_reglas, archivo_reglas):
    df_reglas.fillna('').to_csv(archivo_reglas, index=False)


def cargar_reglas_archivo(ruta):
    """Como cargar_reglas, pero sin datos de ejemplo: si no existe el archivo, devuelve una tabla vacía."""
    if os.path.exists(ruta):
        df = pd.read_csv(ruta).fillna('')
        for col in COLUMNAS_CSV_REGLAS:
            if col not in df.columns:
                df[col] = ''
        return df[COLUMNAS_CSV_REGLAS]
    return pd.DataFrame(columns=COLUMNAS_CSV_REGLAS)


# --- PERFILES DE REGLAS (distintos juegos de reglas guardables, ej. por tipo de reporte) ---

def carpeta_perfiles(perfiles_base_dir):
    ruta = os.path.join(perfiles_base_dir, "perfiles_reglas")
    os.makedirs(ruta, exist_ok=True)
    return ruta


def ruta_perfil(perfiles_base_dir, nombre):
    return os.path.join(carpeta_perfiles(perfiles_base_dir), f"{nombre}.csv")


def listar_perfiles(perfiles_base_dir, archivo_legacy=None):
    carpeta = carpeta_perfiles(perfiles_base_dir)
    nombres = sorted(f[:-4] for f in os.listdir(carpeta) if f.lower().endswith(".csv"))
    if not nombres:
        # Primer uso: migra el CSV de la versión web si existe, o arranca con un perfil de ejemplo.
        if archivo_legacy and os.path.exists(archivo_legacy):
            df_inicial = cargar_reglas(archivo_legacy)
        else:
            df_inicial = reglas_por_defecto()
        guardar_reglas(df_inicial, ruta_perfil(perfiles_base_dir, "Default"))
        nombres = ["Default"]
    return nombres


def cargar_perfil(perfiles_base_dir, nombre):
    return cargar_reglas_archivo(ruta_perfil(perfiles_base_dir, nombre))


def guardar_perfil(perfiles_base_dir, nombre, df_reglas):
    guardar_reglas(df_reglas, ruta_perfil(perfiles_base_dir, nombre))


def eliminar_perfil(perfiles_base_dir, nombre):
    ruta = ruta_perfil(perfiles_base_dir, nombre)
    if os.path.exists(ruta):
        os.remove(ruta)


def nombre_perfil_valido(nombre):
    if not nombre or not nombre.strip():
        return False
    return not any(c in nombre for c in r'\/:*?"<>|')


def detectar_valores(ruta_excel):
    """Lee el Excel del ERP y devuelve los valores únicos por columna, para autocompletar las reglas."""
    df = pd.read_excel(ruta_excel, sheet_name='Sheet2', header=1)
    df.columns = df.columns.str.strip().str.replace('\n', ' ')
    df = df.dropna(subset=['Sucursal']).copy()

    valores = {}
    for col in COLUMNAS_REGLA:
        if col in df.columns:
            valores[col] = sorted({v.strip() for v in df[col].dropna().astype(str) if v.strip()})
        else:
            valores[col] = []

    if 'Razon Social Empresa - Division' in df.columns:
        rs = df['Razon Social Empresa - Division'].astype(str)
        rs = rs.str.replace(r'\s*-\s*Empresa', '', regex=True, case=False)
        rs = rs.str.replace(r'\s*-\s*Prueba', '', regex=True, case=False).str.strip()
        valores['Razón Social'] = sorted({v for v in rs.unique() if v and v.lower() != 'nan'})
    else:
        valores['Razón Social'] = []

    return valores


COLUMNAS_ACUMULADO = ['Mes_Año', 'Mes_Format', 'Sucursal', 'Razón Social', COL_SUBTOTAL, COL_IVA, COL_TOTAL]


def reiniciar_acumulado(archivo_acumulado):
    """Borra el histórico acumulado (para empezar de cero, ej. un año nuevo)."""
    if os.path.exists(archivo_acumulado):
        os.remove(archivo_acumulado)


# Anti-inyección de fórmulas en Excel: un texto de los datos (Razón Social, Sucursal) que
# empiece con = + - @ o un control se interpretaría como fórmula al abrir el archivo. Se le
# antepone una comilla simple para que Excel lo trate como texto literal.
def celda_segura(v):
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r", "\n"):
        return "'" + v
    return v


def generar_reporte(ruta_excel, df_reglas, archivo_acumulado="datos_acumulados.csv", log=print):
    """Procesa el Excel del ERP, lo suma al histórico acumulado, y devuelve (workbook, nombre_archivo_sugerido)
    con TODOS los meses acumulados hasta la fecha (no solo los del archivo recién subido)."""

    # LECTURA Y FILTRADO INICIAL
    df = pd.read_excel(ruta_excel, sheet_name='Sheet2', header=1)
    df.columns = df.columns.str.strip().str.replace('\n', ' ')
    df = df.dropna(subset=['Sucursal']).copy()

    for col in [COL_TOTAL, COL_SUBTOTAL, COL_IVA]:
        df[col] = pd.to_numeric(df[col], errors='coerce').fillna(0)

    # FECHAS
    fechas_limpias = pd.to_datetime(df['Código Fecha Documento'], errors='coerce')
    df = df.loc[fechas_limpias.notna()].copy()
    df['Mes_Año'] = pd.to_datetime(df['Código Fecha Documento']).dt.to_period('M').astype(str)
    df['Mes_Format'] = df['Mes_Año'].apply(formatear_mes_es)

    # LIMPIEZA INTELIGENTE DE RAZÓN SOCIAL
    df['Razón Social'] = df['Razon Social Empresa - Division'].astype(str)
    df['Razón Social'] = df['Razón Social'].str.replace(r'\s*-\s*Empresa', '', regex=True, case=False)
    df['Razón Social'] = df['Razón Social'].str.replace(r'\s*-\s*Prueba', '', regex=True, case=False)
    df['Razón Social'] = df['Razón Social'].str.strip()
    df['Razón Social'] = df['Razón Social'].replace({
        'Bekir I SRL': 'Bekir 1 SRL',
        'Dunia Importer SRL': 'Dunia Import SRL'
    })

    # APLICAR REGLAS
    for _, row in df_reglas.iterrows():
        suc_ori, col, texto, nueva_suc, nueva_rs = map(
            lambda x: str(x).strip(),
            [row.get('Sucursal Origen (Opcional)', ''), row['Columna a Buscar'], row['Contiene el Texto'],
             row['Asignar a Sucursal'], row.get('Asignar a Razón Social', '')]
        )
        if col in df.columns and texto:
            mask = df[col].astype(str).str.contains(texto, case=False, na=False, regex=False)
            if suc_ori:
                mask &= (df['Sucursal'].astype(str).str.upper() == suc_ori.upper())
            df.loc[mask, 'Sucursal'] = nueva_suc
            if nueva_rs:
                df.loc[mask, 'Razón Social'] = nueva_rs

    # EXCLUIR Y REPASO FINAL
    df = df[df['Sucursal'] != 'EXCLUIR']
    df['Razón Social'] = df['Razón Social'].str.replace('Dunia Importer SRL', 'Dunia Import SRL', regex=False)

    log("¡Datos procesados! Actualizando histórico acumulado...")

    # --- ACTUALIZAR EL HISTÓRICO ACUMULADO (resumen por Sucursal y Mes) ---
    # Esto es lo que permite que, mes a mes, se le vaya sumando el mes nuevo sin perder los anteriores,
    # aunque el Excel que subís cada vez traiga solo los datos de ese mes.
    def moda(serie):
        return serie.value_counts().index[0] if len(serie) > 0 else ''

    resumen_nuevo = df.groupby(['Mes_Año', 'Mes_Format', 'Sucursal']).agg(**{
        COL_SUBTOTAL: (COL_SUBTOTAL, 'sum'),
        COL_IVA: (COL_IVA, 'sum'),
        COL_TOTAL: (COL_TOTAL, 'sum'),
        'Razón Social': ('Razón Social', moda),
    }).reset_index()[COLUMNAS_ACUMULADO]

    if os.path.exists(archivo_acumulado):
        acumulado = pd.read_csv(archivo_acumulado, dtype={'Mes_Año': str})
    else:
        acumulado = pd.DataFrame(columns=COLUMNAS_ACUMULADO)

    meses_nuevos = resumen_nuevo['Mes_Año'].unique()
    acumulado = acumulado[~acumulado['Mes_Año'].isin(meses_nuevos)]
    acumulado = pd.concat([acumulado, resumen_nuevo], ignore_index=True)
    acumulado.to_csv(archivo_acumulado, index=False)

    log(f"Histórico acumulado: {acumulado['Mes_Año'].nunique()} mes(es) en total. Generando Excel...")

    # --- GENERACIÓN DE EXCEL (con TODO el histórico acumulado, no solo el archivo recién subido) ---
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws_graficos = wb.create_sheet("Análisis")
    ws_graficos.sheet_view.showGridLines = False
    ws_data = wb.create_sheet("Data_Oculta")
    ws_data.sheet_state = 'hidden'

    df_tot_mes = acumulado.groupby(['Mes_Año', 'Mes_Format'])[COL_TOTAL].sum().reset_index().rename(columns={COL_TOTAL: 'Total Empresa'})
    df_rs_mes = acumulado.groupby(['Mes_Format', 'Razón Social'])[COL_TOTAL].sum().unstack().fillna(0)
    df_loc_mes = acumulado.groupby(['Mes_Format', 'Sucursal'])[COL_TOTAL].sum().unstack().fillna(0)

    df_charts = df_tot_mes.set_index('Mes_Format').drop(columns=['Mes_Año']).join(df_rs_mes).join(df_loc_mes).reset_index()
    ws_data.append([celda_segura(x) for x in df_charts.columns.tolist()])
    for _, r in df_charts.iterrows():
        ws_data.append([celda_segura(x) for x in r.tolist()])

    def dibujar_grafico(titulo, min_col, max_col, dest_cell, tipo="clustered", tendencia_unica=False):
        if max_col < min_col:
            return
        c = BarChart()
        c.type = "col"
        c.grouping = tipo
        if tipo == "stacked":
            c.overlap = 100

        c.title = titulo
        c.style = 13
        c.width, c.height = 30, 14
        c.y_axis.title, c.y_axis.number_format, c.x_axis.title = 'Ventas ($)', '#,##0', 'Mes'
        c.legend.position = 'b'

        data = Reference(ws_data, min_col=min_col, min_row=1, max_col=max_col, max_row=ws_data.max_row)
        cats = Reference(ws_data, min_col=1, min_row=2, max_row=ws_data.max_row)
        c.add_data(data, titles_from_data=True)
        c.set_categories(cats)

        colores = ["1F4E78", "ED7D31", "7030A0", "C00000", "00B0F0", "FFC000", "92D050", "002060", "FF00FF"]
        for i, s in enumerate(c.series):
            s.graphicalProperties.solidFill = colores[i % len(colores)]

        if tendencia_unica and len(c.series) > 0:
            c.series[0].trendline = Trendline(trendlineType='movingAvg', period=2)

        ws_graficos.add_chart(c, dest_cell)

    ws_graficos['B2'] = "ANÁLISIS DE CRECIMIENTO TEMPORAL"
    ws_graficos['B2'].font = Font(size=20, bold=True, color="1F4E78")

    dibujar_grafico("1. Crecimiento Total Empresa", 2, 2, "B4", tipo="clustered")
    dibujar_grafico("2. Crecimiento por Razón Social (Colores separados)", 3, 2 + len(df_rs_mes.columns), "Q4", tipo="clustered", tendencia_unica=True)
    dibujar_grafico("3. Crecimiento de TODAS las Sucursales y Canales (Composición por Mes)", 3 + len(df_rs_mes.columns), ws_data.max_column, "B32", tipo="stacked")

    meses_ordenados = sorted(acumulado['Mes_Año'].unique())

    for mes in meses_ordenados:
        datos_mes = acumulado[acumulado['Mes_Año'] == mes]
        nombre_solapa = datos_mes['Mes_Format'].iloc[0]
        ws_mes = wb.create_sheet(nombre_solapa)
        ws_mes.sheet_view.showGridLines = False

        rk_locales = datos_mes[['Sucursal', 'Razón Social', COL_SUBTOTAL, COL_IVA, COL_TOTAL]].copy()
        rk_locales = rk_locales.sort_values(by=COL_TOTAL, ascending=False).reset_index(drop=True)
        rk_locales.index += 1
        rk_locales = rk_locales.reset_index().rename(columns={'index': 'Rank Actual'})

        rk_empresas = datos_mes.groupby('Razón Social')[[COL_SUBTOTAL, COL_IVA, COL_TOTAL]].sum().reset_index()
        rk_empresas = rk_empresas.sort_values(by=COL_TOTAL, ascending=False).reset_index(drop=True)

        mes_dt = pd.to_datetime(mes + '-01')
        mes_ant_str = (mes_dt - pd.DateOffset(months=1)).strftime('%Y-%m')

        datos_mes_ant = acumulado[acumulado['Mes_Año'] == mes_ant_str]
        if len(datos_mes_ant) > 0:
            rk_ant = datos_mes_ant.sort_values(by=COL_TOTAL, ascending=False).reset_index(drop=True)
            rk_ant.index += 1
            dict_rank_ant = dict(zip(rk_ant['Sucursal'], rk_ant.index))
        else:
            dict_rank_ant = {}
        rk_locales['Rank Anterior'] = rk_locales['Sucursal'].map(dict_rank_ant).fillna(rk_locales['Rank Actual'])
        rk_locales['Diff Num'] = rk_locales['Rank Anterior'] - rk_locales['Rank Actual']
        rk_locales['Diff Rank'] = rk_locales['Diff Num'].apply(lambda x: f"+{int(x)}" if x > 0 else (f"{int(x)}" if x < 0 else "="))

        F_HEAD, BG_DARK = Font(color="FFFFFF", bold=True), PatternFill(start_color="0B3D59", end_color="0B3D59", fill_type="solid")
        C_AL, L_AL, R_AL = Alignment(horizontal="center", vertical="center"), Alignment(horizontal="left", vertical="center"), Alignment(horizontal="right", vertical="center")
        THIN_B = Border(left=Side(style='thin', color="000000"), right=Side(style='thin', color="000000"), top=Side(style='thin', color="000000"), bottom=Side(style='thin', color="000000"))

        ws_mes.merge_cells('A2:K2')
        ws_mes['A2'] = f"TOP POR LOCALES - {nombre_solapa.upper()}"
        ws_mes['A2'].fill, ws_mes['A2'].font, ws_mes['A2'].alignment = BG_DARK, Font(color="FFFFFF", bold=True, size=12), C_AL

        headers = ['Diff Rank', 'Rank', 'Anterior', 'Razón Social', 'Local', 'ARS', 'Subtotal', 'ARS', 'IVA', 'ARS', 'Total']
        for c_i, h in enumerate(headers, 1):
            c = ws_mes.cell(row=3, column=c_i, value=h)
            c.fill, c.font, c.alignment, c.border = BG_DARK, F_HEAD, C_AL, THIN_B

        rn = 4
        for _, r in rk_locales.iterrows():
            c_d = ws_mes.cell(row=rn, column=1, value=r['Diff Rank'])
            c_d.alignment = C_AL
            c_d.font = Font(color="00B050", bold=True) if '+' in str(r['Diff Rank']) else (Font(color="C00000", bold=True) if '-' in str(r['Diff Rank']) else Font(bold=True))

            ws_mes.cell(row=rn, column=2, value=r['Rank Actual']).alignment = C_AL
            ws_mes.cell(row=rn, column=3, value=r['Rank Anterior']).alignment = C_AL
            ws_mes.cell(row=rn, column=4, value=celda_segura(str(r['Razón Social']).split()[0] if str(r['Razón Social']) != 'nan' else 'N/A')).alignment = C_AL
            ws_mes.cell(row=rn, column=5, value=celda_segura(r['Sucursal'])).alignment = L_AL

            for i, v in enumerate([r[COL_SUBTOTAL], r[COL_IVA], r[COL_TOTAL]]):
                ws_mes.cell(row=rn, column=6 + (i * 2), value="ARS").alignment = L_AL
                cx = ws_mes.cell(row=rn, column=7 + (i * 2), value=v)
                cx.number_format, cx.alignment = '#,##0.00', R_AL
            for col in range(1, 12):
                ws_mes.cell(row=rn, column=col).border = THIN_B
            rn += 1

        ws_mes.merge_cells(start_row=rn, start_column=1, end_row=rn, end_column=5)
        ws_mes.cell(row=rn, column=1, value="TOTAL").alignment, ws_mes.cell(row=rn, column=1).font = C_AL, Font(bold=True)
        for i, v in enumerate([rk_locales[COL_SUBTOTAL].sum(), rk_locales[COL_IVA].sum(), rk_locales[COL_TOTAL].sum()]):
            ws_mes.cell(row=rn, column=6 + (i * 2), value="ARS").alignment = L_AL
            cx = ws_mes.cell(row=rn, column=7 + (i * 2), value=v)
            cx.number_format, cx.font = '#,##0.00', Font(bold=True)
        for col in range(1, 12):
            ws_mes.cell(row=rn, column=col).border = THIN_B

        color_scale = ColorScaleRule(start_type='min', start_color='F8696B', mid_type='percentile', mid_value=50, mid_color='FFEB84', end_type='max', end_color='63BE7B')
        for c_l in ['G', 'I', 'K']:
            ws_mes.conditional_formatting.add(f"{c_l}4:{c_l}{rn-1}", color_scale)

        rn += 3
        ws_mes.merge_cells(start_row=rn, start_column=4, end_row=rn, end_column=11)
        ws_mes.cell(row=rn, column=4, value="TOP POR EMPRESAS").fill, ws_mes.cell(row=rn, column=4).font, ws_mes.cell(row=rn, column=4).alignment = BG_DARK, F_HEAD, C_AL
        rn += 1
        ws_mes.merge_cells(start_row=rn, start_column=4, end_row=rn, end_column=5)
        ws_mes.cell(row=rn, column=4, value="RAZON SOCIAL").fill, ws_mes.cell(row=rn, column=4).font, ws_mes.cell(row=rn, column=4).alignment = BG_DARK, F_HEAD, C_AL
        for idx, text in enumerate(["SUBTOTAL", "IVA", "TOTAL"]):
            ws_mes.merge_cells(start_row=rn, start_column=6 + (idx * 2), end_row=rn, end_column=7 + (idx * 2))
            c = ws_mes.cell(row=rn, column=6 + (idx * 2), value=text)
            c.fill, c.font, c.alignment = BG_DARK, F_HEAD, C_AL

        rn += 1
        for _, r in rk_empresas.iterrows():
            ws_mes.merge_cells(start_row=rn, start_column=4, end_row=rn, end_column=5)
            ws_mes.cell(row=rn, column=4, value=celda_segura(r['Razón Social'])).alignment = C_AL
            for i, v in enumerate([r[COL_SUBTOTAL], r[COL_IVA], r[COL_TOTAL]]):
                ws_mes.cell(row=rn, column=6 + (i * 2), value="ARS").alignment = L_AL
                cx = ws_mes.cell(row=rn, column=7 + (i * 2), value=v)
                cx.number_format = '#,##0.00'
            for col in range(4, 12):
                ws_mes.cell(row=rn, column=col).border = THIN_B
            rn += 1

        ws_mes.merge_cells(start_row=rn, start_column=4, end_row=rn, end_column=5)
        ws_mes.cell(row=rn, column=4, value="TOTAL").alignment, ws_mes.cell(row=rn, column=4).font = C_AL, Font(bold=True)
        for i, v in enumerate([rk_empresas[COL_SUBTOTAL].sum(), rk_empresas[COL_IVA].sum(), rk_empresas[COL_TOTAL].sum()]):
            ws_mes.cell(row=rn, column=6 + (i * 2), value="ARS").alignment = L_AL
            cx = ws_mes.cell(row=rn, column=7 + (i * 2), value=v)
            cx.number_format, cx.font = '#,##0.00', Font(bold=True)
        for col in range(4, 12):
            ws_mes.cell(row=rn, column=col).border = THIN_B

        ws_mes.column_dimensions['A'].width, ws_mes.column_dimensions['B'].width, ws_mes.column_dimensions['C'].width = 8, 8, 8
        ws_mes.column_dimensions['D'].width, ws_mes.column_dimensions['E'].width = 15, 25
        for c in ['F', 'H', 'J']:
            ws_mes.column_dimensions[c].width = 5
        for c in ['G', 'I', 'K']:
            ws_mes.column_dimensions[c].width = 17

        pie_loc = PieChart()
        pie_loc.add_data(Reference(ws_mes, min_col=11, min_row=3, max_row=3 + len(rk_locales)), titles_from_data=True)
        pie_loc.set_categories(Reference(ws_mes, min_col=5, min_row=4, max_row=3 + len(rk_locales)))
        pie_loc.title = "LOCALES"
        pie_loc.dataLabels = DataLabelList()
        pie_loc.dataLabels.showPercent, pie_loc.dataLabels.showCatName = True, True
        pie_loc.width, pie_loc.height = 15, 12
        ws_mes.add_chart(pie_loc, "M3")

        pie_emp = PieChart()
        start_emp_row = rn - len(rk_empresas)
        pie_emp.add_data(Reference(ws_mes, min_col=11, min_row=start_emp_row - 1, max_row=rn - 1), titles_from_data=True)
        pie_emp.set_categories(Reference(ws_mes, min_col=4, min_row=start_emp_row, max_row=rn - 1))
        pie_emp.title = "RAZON SOCIAL"
        pie_emp.dataLabels = DataLabelList()
        pie_emp.dataLabels.showPercent, pie_emp.dataLabels.showCatName = True, True
        pie_emp.width, pie_emp.height = 15, 12
        ws_mes.add_chart(pie_emp, "U3")

    if len(meses_ordenados) == 1:
        nom_arch = f"Reporte_{formatear_mes_es(meses_ordenados[0]).replace(' ', '_')}.xlsx"
    else:
        nom_arch = f"Reporte_Acumulado_{formatear_mes_es(meses_ordenados[0]).replace(' ', '_')}_a_{formatear_mes_es(meses_ordenados[-1]).replace(' ', '_')}.xlsx"

    return wb, nom_arch
